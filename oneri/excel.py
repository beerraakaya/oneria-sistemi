"""Öneri Excel'ini okuyup her satırı sade bir kayda çevirir.

Kişi adı sütunları hiç okunmaz; yapay zekâya yalnızca öneri bilgileri gider.
"""

import hashlib
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ONERI = "Öneri"
ONERI_DEGIL = "Öneri Değil"
# Onay Durumu'na göre Durum sütununa yazılan başlangıç değeri.
# Tamamlandı ve Uygulanamaz, süreç ilerledikçe ekip tarafından girilir.
BASLANGIC_DURUMU = {ONERI: "Devam Ediyor", ONERI_DEGIL: "Red Edildi"}

# Alan adı -> Excel başlığı. Uzun başlıklar baştaki kısımlarıyla eşleşir.
BASLIKLAR = {
    "tarih": "Tarih",
    "fabrika": "Önerinin Uygulanacağı Yeri",
    "bolum": "Bölüm",
    "konu": "Öneri Konusu",
    "mevcut_durum": "Mevcut Durum",
    "onerilen_durum": "Önerilen Durum",
    "degerlendirme": "Değerlendirme",
    "onay_durumu": "Onay Durumu",
    "durum": "Durum",
}


class ExcelHatasi(ValueError):
    """Excel dosyası bulunamadı ya da beklenen yapıda değil."""


@dataclass(frozen=True)
class Oneri:
    satir: int  # Excel'deki satır numarası
    tarih: date | None
    fabrika: str
    bolum: str
    konu: str
    mevcut_durum: str
    onerilen_durum: str
    degerlendirme: str
    onay_durumu: str
    durum: str

    @property
    def bekliyor(self) -> bool:
        """Ekip bu öneriye henüz hiçbir şey yazmamışsa True."""
        return not (self.degerlendirme or self.onay_durumu or self.durum)

    @property
    def anahtar(self) -> str:
        """Satır numarası değişse (sıralama, silme) de öneriyi tanımaya yarayan kimlik."""
        tarih = self.tarih.isoformat() if self.tarih else ""
        parcalar = (tarih, self.konu, self.mevcut_durum, self.onerilen_durum)
        return hashlib.sha256("\x1f".join(parcalar).encode("utf-8")).hexdigest()

    def arama_metni(self) -> str:
        """Benzer önerileri bulurken karşılaştırılan metin."""
        return "\n".join((self.konu, self.mevcut_durum, self.onerilen_durum))


def onerileri_oku(yol: Path, sayfa_adi: str) -> list[Oneri]:
    kitap = _ac(yol, sayfa_adi)
    try:
        satirlar = kitap[sayfa_adi].iter_rows(values_only=True)
        sutun = _sutunlar(next(satirlar, ()))

        oneriler = []
        for satir_no, hucreler in enumerate(satirlar, start=2):
            deger = {alan: _hucre(hucreler, i) for alan, i in sutun.items()}
            if not sade_metin(deger["fabrika"]) and not sade_metin(deger["konu"]):
                continue  # boş ya da öneri içermeyen satır
            oneriler.append(
                Oneri(
                    satir=satir_no,
                    tarih=_tarih(deger["tarih"]),
                    fabrika=sade_metin(deger["fabrika"]),
                    bolum=sade_metin(deger["bolum"]),
                    konu=sade_metin(deger["konu"]),
                    mevcut_durum=sade_metin(deger["mevcut_durum"]),
                    onerilen_durum=sade_metin(deger["onerilen_durum"]),
                    degerlendirme=sade_metin(deger["degerlendirme"]),
                    onay_durumu=_onay(deger["onay_durumu"]),
                    durum=sade_metin(deger["durum"]),
                )
            )
        return oneriler
    finally:
        kitap.close()


def sutun_harfleri(yol: Path, sayfa_adi: str) -> dict[str, str]:
    """Alan adı -> Excel sütun harfi, örn. {"degerlendirme": "O", ...}."""
    kitap = _ac(yol, sayfa_adi)
    try:
        baslik_satiri = next(kitap[sayfa_adi].iter_rows(values_only=True), ())
        return {alan: get_column_letter(i + 1) for alan, i in _sutunlar(baslik_satiri).items()}
    finally:
        kitap.close()


def _ac(yol: Path, sayfa_adi: str):
    if not yol.exists():
        raise ExcelHatasi(f"Excel dosyası bulunamadı: {yol}")
    kitap = openpyxl.load_workbook(yol, read_only=True, data_only=True)
    if sayfa_adi not in kitap.sheetnames:
        kitap.close()
        raise ExcelHatasi(
            f"'{sayfa_adi}' sayfası bulunamadı. Dosyadaki sayfalar: {', '.join(kitap.sheetnames)}"
        )
    return kitap


def _sutunlar(baslik_satiri) -> dict[str, int]:
    basliklar = [sade_metin(baslik) for baslik in baslik_satiri]
    return {alan: _sutun_bul(basliklar, baslik) for alan, baslik in BASLIKLAR.items()}


def _sutun_bul(basliklar: list[str], aranan: str) -> int:
    if aranan in basliklar:
        return basliklar.index(aranan)
    for i, baslik in enumerate(basliklar):
        if baslik.startswith(aranan):
            return i
    raise ExcelHatasi(f"Excel'de '{aranan}' sütunu bulunamadı.")


def _hucre(hucreler: tuple, i: int):
    return hucreler[i] if i < len(hucreler) else None


def sade_metin(deger) -> str:
    if deger is None:
        return ""
    return re.sub(r"\s+", " ", str(deger)).strip()


def _onay(deger) -> str:
    """Sonunda boşluk kalmış "Öneri " gibi yazım farklarını tek biçime getirir."""
    metin = sade_metin(deger)
    for secenek in BASLANGIC_DURUMU:
        if metin.casefold() == secenek.casefold():
            return secenek
    return metin


def _tarih(deger) -> date | None:
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    if isinstance(deger, str):
        # Otomatik eklenen satırlarda tarih "2025-01-07T00:00:00+00:27" gibi metin olarak geliyor.
        try:
            return date.fromisoformat(deger.strip()[:10])
        except ValueError:
            return None
    return None
