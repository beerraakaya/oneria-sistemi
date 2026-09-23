"""Kör test: cevabı bilinen önerileri, cevapları gizleyerek yapay zekâya değerlendirtir."""

import time
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .ayarlar import Ayarlar
from .degerlendirici import Degerlendirici, Taslak
from .excel import ONERI, ONERI_DEGIL, Oneri
from .ollama import GecersizCevap, OllamaHatasi


@dataclass(frozen=True)
class Sonuc:
    oneri: Oneri
    taslak: Taslak | None  # None ise model geçerli bir cevap veremedi
    hata: str
    sure: float


@dataclass(frozen=True)
class Ozet:
    toplam: int
    cevaplanan: int
    ayni_karar: int
    # (ekibin kararı, yapay zekânın kararı) -> adet
    karsilastirma: Counter
    ortalama_sure: float

    @property
    def oran(self) -> float:
        return 100 * self.ayni_karar / self.cevaplanan if self.cevaplanan else 0.0


def kor_test_calistir(
    degerlendirici: Degerlendirici,
    oneriler: list[Oneri],
    ilerleme: Callable[[str], None] = print,
) -> list[Sonuc]:
    """Her öneri, kendisi örneklerden çıkarılarak değerlendirilir.

    Test yarıda kesilirse (Ctrl+C ya da Ollama hatası) o ana kadarki sonuçlar döner.
    """
    sonuclar = []
    for sira, oneri in enumerate(oneriler, start=1):
        baslangic = time.monotonic()
        try:
            taslak, hata = degerlendirici.degerlendir(oneri, haric_satir=oneri.satir), ""
        except GecersizCevap as h:
            taslak, hata = None, str(h)
        except (OllamaHatasi, KeyboardInterrupt) as h:
            ilerleme(f"Test yarıda kesildi: {str(h) or 'kullanıcı durdurdu'}")
            break
        sonuc = Sonuc(oneri, taslak, hata, time.monotonic() - baslangic)
        sonuclar.append(sonuc)
        if taslak is None:
            durum = "cevap alınamadı"
        elif taslak.onay_durumu == oneri.onay_durumu:
            durum = f"aynı karar ({taslak.onay_durumu})"
        else:
            durum = f"FARKLI karar: ekip {oneri.onay_durumu}, yapay zekâ {taslak.onay_durumu}"
        ilerleme(f"[{sira}/{len(oneriler)}] satır {oneri.satir}: {durum} ({sonuc.sure:.0f} sn)")
    return sonuclar


def ozetle(sonuclar: list[Sonuc]) -> Ozet:
    cevaplanan = [s for s in sonuclar if s.taslak]
    return Ozet(
        toplam=len(sonuclar),
        cevaplanan=len(cevaplanan),
        ayni_karar=sum(s.taslak.onay_durumu == s.oneri.onay_durumu for s in cevaplanan),
        karsilastirma=Counter((s.oneri.onay_durumu, s.taslak.onay_durumu) for s in cevaplanan),
        ortalama_sure=sum(s.sure for s in sonuclar) / len(sonuclar) if sonuclar else 0.0,
    )


_BASLIK_DOLGUSU = PatternFill("solid", fgColor="DCE6F1")
_FARKLI_DOLGUSU = PatternFill("solid", fgColor="F8D7DA")
_KALIN = Font(bold=True)
_SARMALA = Alignment(wrap_text=True, vertical="top")

# (başlık, sütun genişliği)
_SUTUNLAR = [
    ("Satır", 7),
    ("Tarih", 11),
    ("Konu", 22),
    ("Mevcut Durum", 40),
    ("Önerilen Durum", 40),
    ("Ekibin Kararı", 13),
    ("Yapay Zekânın Kararı", 13),
    ("Aynı mı?", 9),
    ("Ekibin Değerlendirmesi", 55),
    ("Yapay Zekânın Değerlendirmesi", 55),
    ("Metin Puanı (1-5)", 11),
    ("Not", 30),
    ("Süre (sn)", 9),
    ("Hata", 30),
]


def rapor_yaz(sonuclar: list[Sonuc], yol: Path, ayarlar: Ayarlar) -> None:
    kitap = Workbook()
    _ozet_sayfasi(kitap.active, ozetle(sonuclar), ayarlar)
    _karsilastirma_sayfasi(kitap.create_sheet("Karşılaştırma"), sonuclar)
    yol.parent.mkdir(parents=True, exist_ok=True)
    kitap.save(yol)


def _ozet_sayfasi(sayfa, ozet: Ozet, ayarlar: Ayarlar) -> None:
    sayfa.title = "Özet"
    satirlar = [
        ("Kör test raporu", ""),
        ("Tarih", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Dil modeli", ayarlar.dil_modeli),
        ("Gömme modeli", ayarlar.gomme_modeli),
        ("", ""),
        ("Test edilen öneri", ozet.toplam),
        ("Cevap alınan", ozet.cevaplanan),
        ("Ekiple aynı karar", f"{ozet.ayni_karar} / {ozet.cevaplanan} (%{ozet.oran:.0f})"),
        ("", ""),
        ("Ekibin kararı -> Yapay zekânın kararı", "Adet"),
    ]
    tablo_basligi = len(satirlar)
    for ekip in (ONERI, ONERI_DEGIL):
        for yapay_zeka in (ONERI, ONERI_DEGIL):
            satirlar.append((f"{ekip} -> {yapay_zeka}", ozet.karsilastirma[(ekip, yapay_zeka)]))
    satirlar += [
        ("", ""),
        ("Öneri başına ortalama süre (sn)", round(ozet.ortalama_sure, 1)),
        ("", ""),
        (
            "Not",
            "Karşılaştırma sayfasındaki 'Metin Puanı' sütununa 1-5 arası puan vererek "
            "yapay zekânın metinlerini de değerlendirebilirsiniz.",
        ),
    ]
    for satir in satirlar:
        sayfa.append(satir)
    sayfa["A1"].font = Font(bold=True, size=14)
    for sutun in (1, 2):
        sayfa.cell(row=tablo_basligi, column=sutun).font = _KALIN
    sayfa.column_dimensions["A"].width = 38
    sayfa.column_dimensions["B"].width = 60
    sayfa.cell(row=len(satirlar), column=2).alignment = _SARMALA


def _karsilastirma_sayfasi(sayfa, sonuclar: list[Sonuc]) -> None:
    sayfa.append([baslik for baslik, _ in _SUTUNLAR])
    for sutun, (_, genislik) in enumerate(_SUTUNLAR, start=1):
        hucre = sayfa.cell(row=1, column=sutun)
        hucre.font = _KALIN
        hucre.fill = _BASLIK_DOLGUSU
        hucre.alignment = _SARMALA
        sayfa.column_dimensions[hucre.column_letter].width = genislik

    for s in sonuclar:
        ayni = s.taslak is not None and s.taslak.onay_durumu == s.oneri.onay_durumu
        sayfa.append(
            [
                s.oneri.satir,
                s.oneri.tarih.strftime("%d.%m.%Y") if s.oneri.tarih else "",
                s.oneri.konu,
                s.oneri.mevcut_durum,
                s.oneri.onerilen_durum,
                s.oneri.onay_durumu,
                s.taslak.onay_durumu if s.taslak else "",
                ("Evet" if ayni else "Hayır") if s.taslak else "",
                s.oneri.degerlendirme,
                s.taslak.degerlendirme if s.taslak else "",
                None,
                None,
                round(s.sure, 1),
                s.hata,
            ]
        )
        for hucre in sayfa[sayfa.max_row]:
            hucre.alignment = _SARMALA
        if s.taslak and not ayni:
            for sutun in (6, 7, 8):
                sayfa.cell(row=sayfa.max_row, column=sutun).fill = _FARKLI_DOLGUSU

    puan = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    puan.error = "1 ile 5 arasında bir tam sayı girin."
    sayfa.add_data_validation(puan)
    if sonuclar:
        puan.add(f"K2:K{len(sonuclar) + 1}")
    sayfa.freeze_panes = "A2"
