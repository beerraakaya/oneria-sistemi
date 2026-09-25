"""Öneri Excel'ine erişim yolları:

- SharePointExcel: SharePoint'teki dosya, IT'nin verdiği uygulama izniyle (Microsoft Graph).
- ExcelUygulamasi: SharePoint'teki dosya, bilgisayardaki Excel uygulaması ve oturum açmış
  kullanıcının hesabıyla (IT izni gerekmez; yalnızca Windows).
- YerelExcel: bilgisayardaki bir kopya (denemek için).

Program Excel'e yalnızca hücre değeri ve dolgu rengi yazar; satır eklemez, silmez, sıralamaz.
"""

import base64
import time
from collections.abc import Callable
from pathlib import Path
from typing import Protocol
from urllib.parse import quote, unquote, urlsplit

import openpyxl
import requests
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from .excel import ExcelHatasi


class ExcelKaynagi(Protocol):
    ad: str

    def indir(self, hedef: Path) -> Path:
        """Dosyanın güncel hâlini okunabilir bir yola getirir ve o yolu döndürür."""

    def satir_oku(self, sayfa: str, satir: int, sutun_sayisi: int) -> list:
        """Bir satırın ilk `sutun_sayisi` hücresinin şu anki değerleri."""

    def yaz(self, sayfa: str, hucreler: dict[str, str]) -> None:
        """{"O567": "metin", ...} hücrelerine yazar."""

    def boya(self, sayfa: str, adresler: list[str], renk: str | None) -> None:
        """Hücrelerin dolgu rengini değiştirir; renk None ise dolguyu kaldırır."""

    def kapat(self) -> None:
        """Açık bağlantıları ve uygulamaları kapatır."""


class YerelExcel:
    """Bilgisayardaki bir Excel dosyası. Canlıya geçmeden önce bir kopyada denemek içindir.

    Dosya Excel'de açıkken yazılamaz. openpyxl kaydederken grafik gibi bazı öğeleri
    koruyamayabileceği için asıl dosyada değil, kopyada kullanılmalıdır.
    """

    def __init__(self, yol: Path):
        self._yol = yol
        self.ad = str(yol)

    def indir(self, hedef: Path) -> Path:
        return self._yol

    def satir_oku(self, sayfa: str, satir: int, sutun_sayisi: int) -> list:
        kitap = openpyxl.load_workbook(self._yol, read_only=True, data_only=True)
        try:
            hucreler = next(
                kitap[sayfa].iter_rows(
                    min_row=satir, max_row=satir, max_col=sutun_sayisi, values_only=True
                ),
                (),
            )
            return list(hucreler) + [None] * (sutun_sayisi - len(hucreler))
        finally:
            kitap.close()

    def yaz(self, sayfa: str, hucreler: dict[str, str]) -> None:
        def uygula(s):
            for adres, deger in hucreler.items():
                s[adres] = deger

        self._degistir(sayfa, uygula)

    def boya(self, sayfa: str, adresler: list[str], renk: str | None) -> None:
        dolgu = (
            PatternFill(fill_type="solid", fgColor=renk.lstrip("#").upper())
            if renk
            else PatternFill(fill_type=None)
        )

        def uygula(s):
            for adres in adresler:
                s[adres].fill = dolgu

        self._degistir(sayfa, uygula)

    def kapat(self) -> None:
        pass

    def _degistir(self, sayfa: str, islem: Callable) -> None:
        kitap = openpyxl.load_workbook(self._yol)
        try:
            islem(kitap[sayfa])
            kitap.save(self._yol)
        except PermissionError as hata:
            raise ExcelHatasi(f"{self._yol} yazılamadı; dosya Excel'de açık olabilir.") from hata
        finally:
            kitap.close()


class GraphHatasi(RuntimeError):
    """SharePoint'e (Microsoft Graph) erişilemedi."""


class SharePointExcel:
    """SharePoint'teki Excel'e, IT'nin Microsoft Entra'da açtığı uygulama kimliğiyle erişir.

    Uygulama izni (application permission) kullanılır; kimse oturum açmadan çalışır.
    Excel başka kişilerde açıkken de yazılabilir.
    """

    def __init__(
        self,
        dosya_adresi: str,
        kiraci: str,
        uygulama: str,
        sir: str,
        *,
        giris_adresi: str = "https://login.microsoftonline.com",
        graph_adresi: str = "https://graph.microsoft.com/v1.0",
        bekle: Callable[[float], None] = time.sleep,
    ):
        self._dosya_adresi = dosya_adresi.strip()
        self._kiraci = kiraci
        self._uygulama = uygulama
        self._sir = sir
        self._giris_adresi = giris_adresi.rstrip("/")
        self._graph_adresi = graph_adresi.rstrip("/")
        self._bekle = bekle
        # Ollama'nın aksine burada şirket proxy'si kullanılmalı; ortam ayarları okunur.
        self._oturum = requests.Session()
        self._jeton = ""
        self._jeton_bitis = 0.0
        self._dosya_yolu = ""  # /drives/{surucu}/items/{oge}
        self.ad = dosya_adresi

    # --- Excel işlemleri ---

    def indir(self, hedef: Path) -> Path:
        cevap = self._istek("GET", f"{self._dosya()}/content")
        hedef.parent.mkdir(parents=True, exist_ok=True)
        hedef.write_bytes(cevap.content)
        return hedef

    def satir_oku(self, sayfa: str, satir: int, sutun_sayisi: int) -> list:
        adres = f"A{satir}:{get_column_letter(sutun_sayisi)}{satir}"
        degerler = self._istek("GET", self._aralik(sayfa, adres)).json().get("values") or [[]]
        hucreler = [None if d == "" else d for d in degerler[0]]
        return hucreler + [None] * (sutun_sayisi - len(hucreler))

    def yaz(self, sayfa: str, hucreler: dict[str, str]) -> None:
        for adres, deger in hucreler.items():
            self._istek("PATCH", self._aralik(sayfa, adres), json={"values": [[deger]]})

    def boya(self, sayfa: str, adresler: list[str], renk: str | None) -> None:
        for adres in adresler:
            if renk:
                self._istek("PATCH", f"{self._aralik(sayfa, adres)}/format/fill", json={"color": renk})
            else:
                self._istek("POST", f"{self._aralik(sayfa, adres)}/format/fill/clear")

    def kapat(self) -> None:
        self._oturum.close()

    # --- Dosyayı bulma ---

    def _aralik(self, sayfa: str, adres: str) -> str:
        return f"{self._dosya()}/workbook/worksheets/{quote(sayfa, safe='')}/range(address='{adres}')"

    def _dosya(self) -> str:
        if not self._dosya_yolu:
            oge = self._paylasim_ogesi() if "/:" in self._dosya_adresi else self._yol_ogesi()
            self._dosya_yolu = f"/drives/{oge['parentReference']['driveId']}/items/{oge['id']}"
        return self._dosya_yolu

    def _paylasim_ogesi(self) -> dict:
        """"Bağlantıyı kopyala" ile alınan paylaşım bağlantısı (…sharepoint.com/:x:/s/…)."""
        kod = base64.urlsafe_b64encode(self._dosya_adresi.encode("utf-8")).decode().rstrip("=")
        return self._istek("GET", f"/shares/u!{kod}/driveItem").json()

    def _yol_ogesi(self) -> dict:
        """Dosyanın tam yolu: https://sirket.sharepoint.com/sites/Site/Shared Documents/…/Dosya.xlsx"""
        parca = urlsplit(self._dosya_adresi)
        if not parca.netloc.endswith(".sharepoint.com"):
            raise GraphHatasi(f"SharePoint adresi anlaşılamadı: {self._dosya_adresi}")
        yol = unquote(parca.path).strip("/").split("/")
        site = "/".join(yol[:2]) if yol[0] in ("sites", "teams") else ""
        site_kimligi = self._istek(
            "GET", f"/sites/{parca.netloc}" + (f":/{quote(site)}" if site else "")
        ).json()["id"]

        tam_adres = f"https://{parca.netloc}/{'/'.join(yol)}"
        suruculer = self._istek("GET", f"/sites/{site_kimligi}/drives").json().get("value", [])
        # Dosyanın bulunduğu belge kitaplığı: adresi dosya adresinin başıyla eşleşen en uzun kitaplık.
        uygun = [
            s for s in suruculer if tam_adres.startswith(unquote(s.get("webUrl", "")).rstrip("/") + "/")
        ]
        if not uygun:
            raise GraphHatasi(f"Dosyanın belge kitaplığı bulunamadı: {self._dosya_adresi}")
        surucu = max(uygun, key=lambda s: len(s["webUrl"]))
        goreli = tam_adres[len(unquote(surucu["webUrl"]).rstrip("/")) + 1 :]
        return self._istek("GET", f"/drives/{surucu['id']}/root:/{quote(goreli)}").json()

    # --- HTTP ---

    def _yetki(self) -> str:
        if time.time() < self._jeton_bitis - 300:
            return self._jeton
        try:
            cevap = self._oturum.post(
                f"{self._giris_adresi}/{self._kiraci}/oauth2/v2.0/token",
                data={
                    "client_id": self._uygulama,
                    "client_secret": self._sir,
                    "scope": "https://graph.microsoft.com/.default",
                    "grant_type": "client_credentials",
                },
                timeout=60,
            )
        except requests.RequestException as hata:
            raise GraphHatasi(f"Microsoft oturum açma sunucusuna bağlanılamadı: {hata}") from hata
        if cevap.status_code != 200:
            raise GraphHatasi(
                "Microsoft'tan erişim izni alınamadı; graph_kiraci, graph_uygulama ve gizli "
                f"anahtarı kontrol edin. ({cevap.status_code}: {_hata_mesaji(cevap)})"
            )
        veri = cevap.json()
        self._jeton = veri["access_token"]
        self._jeton_bitis = time.time() + int(veri.get("expires_in", 3600))
        return self._jeton

    def _istek(self, yontem: str, yol: str, **secenekler) -> requests.Response:
        for deneme in range(4):
            son = deneme == 3
            try:
                cevap = self._oturum.request(
                    yontem,
                    self._graph_adresi + yol,
                    headers={"Authorization": f"Bearer {self._yetki()}"},
                    timeout=120,
                    **secenekler,
                )
            except requests.RequestException as hata:
                if son:
                    raise GraphHatasi(f"SharePoint'e bağlanılamadı: {hata}") from hata
                self._bekle(2**deneme)
                continue

            if cevap.status_code == 401 and deneme == 0:
                self._jeton_bitis = 0  # süresi dolmuş olabilir; yenisini al
                continue
            if cevap.status_code in (429, 500, 502, 503, 504) and not son:
                # SharePoint yoğunken "biraz sonra dene" der; söylediği kadar beklenir.
                self._bekle(min(float(cevap.headers.get("Retry-After") or 2**deneme), 60))
                continue
            if cevap.status_code >= 400:
                raise GraphHatasi(_anlasilir_hata(cevap))
            return cevap
        raise GraphHatasi("SharePoint'e bağlanılamadı.")


def _hata_mesaji(cevap: requests.Response) -> str:
    try:
        veri = cevap.json()
    except ValueError:
        return cevap.text[:200]
    hata = veri.get("error")
    if isinstance(hata, dict):
        return hata.get("message") or hata.get("code") or str(hata)
    return veri.get("error_description") or str(hata or veri)[:200]


def _anlasilir_hata(cevap: requests.Response) -> str:
    mesaj = _hata_mesaji(cevap)
    if cevap.status_code == 403:
        return (
            "Uygulamanın bu dosyaya erişim izni yok (403). IT'den uygulamaya bu SharePoint "
            f"sitesi için yazma izni verilmesini isteyin. ({mesaj})"
        )
    if cevap.status_code == 404:
        return f"SharePoint'te bulunamadı (404); dosya adresini ve sayfa adını kontrol edin. ({mesaj})"
    if cevap.status_code == 423:
        return f"Excel dosyası şu an kilitli (423); bir sonraki çalışmada tekrar denenecek. ({mesaj})"
    return f"SharePoint hatası {cevap.status_code}: {mesaj}"


# Excel'in Interior.ColorIndex için "dolgu yok" değeri (xlNone).
_DOLGU_YOK = -4142
# Excel açılırken ya da arka planda çalışırken komutları "meşgulüm" diye geri çevirir:
# RPC_E_CALL_REJECTED ("Call was rejected by callee") ve RPC_E_SERVERCALL_RETRYLATER.
_EXCEL_MESGUL = {-2147418111, -2147417846}


def _excel_baslat():
    try:
        import win32com.client  # pywin32; yalnızca Windows'ta kurulur
    except ImportError as hata:
        raise ExcelHatasi(
            "Excel yöntemi için pywin32 gerekli ve yalnızca Windows'ta çalışır: "
            "python -m pip install -r requirements.txt"
        ) from hata
    # DispatchEx her seferinde ayrı bir Excel açar; kullanıcının açık Excel'ine dokunulmaz.
    return win32com.client.DispatchEx("Excel.Application")


def _acik_kitabi_bul(adresler: list[str]):
    """Dosya bu bilgisayarda çalışan bir Excel'de zaten açıksa (Excel, kitap) döndürür.

    Aynı hesap aynı SharePoint dosyasını iki ayrı Excel'de açamaz; açıksa ona yazılır.
    """
    try:
        import win32com.client

        excel = win32com.client.GetActiveObject("Excel.Application")
    except Exception:
        return None  # çalışan Excel yok
    aranan = {unquote(a).casefold() for a in adresler}
    try:
        for kitap in excel.Workbooks:
            if unquote(str(kitap.FullName)).casefold() in aranan:
                return excel, kitap
    except Exception:
        return None
    return None


class ExcelUygulamasi:
    """Dosyayı bilgisayardaki Excel uygulamasıyla, oturum açmış kullanıcının hesabıyla açar.

    Bir insanın Excel'de dosyayı açıp hücreye yazması ve kaydetmesiyle aynıdır; IT izni
    gerekmez. Excel ekranda görünmez. Dosya bu bilgisayarda zaten açıksa açık olana yazılır
    ve o pencere kapatılmaz. Her yazmadan sonra kaydedilir; birlikte düzenleme
    sayesinde dosya başkalarında açıkken de yazılabilir ve Excel dosyanın yapısını korur.
    """

    def __init__(
        self,
        dosya_adresi: str,
        *,
        gorunur: bool = False,
        excel_olustur: Callable = _excel_baslat,
        acik_kitap_bul: Callable = _acik_kitabi_bul,
        bekle: Callable[[float], None] = time.sleep,
        mesgul_suresi: float = 120,
    ):
        # "?web=1" gibi ekler Excel'in dosyayı açmasını engelleyebilir.
        self._adres = dosya_adresi.strip().split("?", 1)[0]
        self._gorunur = gorunur
        self._excel_olustur = excel_olustur
        self._acik_kitap_bul = acik_kitap_bul
        self._bekle = bekle
        # False ise dosya zaten açıktı; iş bitince kapatılmaz.
        self._kendimiz_actik = True
        self._mesgul_suresi = mesgul_suresi
        self._excel = None
        self._kitap = None
        self.ad = self._adres

    def indir(self, hedef: Path) -> Path:
        kitap = self._ac()
        hedef.parent.mkdir(parents=True, exist_ok=True)
        hedef.unlink(missing_ok=True)
        try:
            self._tekrarla(lambda: kitap.SaveCopyAs(str(hedef.resolve())))
        except Exception as hata:
            raise ExcelHatasi(f"Dosyanın kopyası alınamadı: {hata}") from hata
        return hedef

    def satir_oku(self, sayfa: str, satir: int, sutun_sayisi: int) -> list:
        tablo = self._sayfa(sayfa)
        adres = f"A{satir}:{get_column_letter(sutun_sayisi)}{satir}"
        degerler = self._tekrarla(lambda: tablo.Range(adres).Value)
        hucreler = list(degerler[0]) if isinstance(degerler, tuple) else [degerler]
        return hucreler + [None] * (sutun_sayisi - len(hucreler))

    def yaz(self, sayfa: str, hucreler: dict[str, str]) -> None:
        tablo = self._sayfa(sayfa)
        for adres, deger in hucreler.items():
            self._tekrarla(lambda: setattr(tablo.Range(adres), "Value", deger))
        self._kaydet()

    def boya(self, sayfa: str, adresler: list[str], renk: str | None) -> None:
        tablo = self._sayfa(sayfa)
        if renk:
            kod = renk.lstrip("#")
            # Excel renkleri ters sırada (mavi, yeşil, kırmızı) tutar.
            ozellik, deger = "Color", int(kod[4:6] + kod[2:4] + kod[0:2], 16)
        else:
            ozellik, deger = "ColorIndex", _DOLGU_YOK
        for adres in adresler:
            self._tekrarla(lambda: setattr(tablo.Range(adres).Interior, ozellik, deger))
        self._kaydet()

    def kapat(self) -> None:
        if not self._kendimiz_actik:
            # Kullanıcının açık Excel'i: yazılanlar zaten kaydedildi, pencereye dokunulmaz.
            self._kitap = self._excel = None
            return
        if self._kitap is not None:
            try:
                kitap = self._kitap
                self._tekrarla(lambda: kitap.Close(SaveChanges=True), sure=10)
            except Exception:
                pass
            self._kitap = None
        if self._excel is not None:
            try:
                excel = self._excel
                self._tekrarla(lambda: excel.Quit(), sure=10)
            except Exception:
                pass
            self._excel = None

    def _ac(self):
        if self._kitap is not None:
            return self._kitap
        bulunan = self._acik_kitap_bul(self._adres_adaylari())
        if bulunan is not None:
            self._excel, self._kitap = bulunan
            self._kendimiz_actik = False
            if self._tekrarla(lambda: self._kitap.ReadOnly):
                self._kitap = self._excel = None
                raise ExcelHatasi(
                    "Dosya bu bilgisayarda salt okunur olarak açık. Excel'deki dosyayı kapatın "
                    "ya da düzenlemeyi etkinleştirin."
                )
            return self._kitap
        self._kendimiz_actik = True
        self._excel = self._excel_olustur()
        try:
            excel = self._excel
            self._tekrarla(lambda: setattr(excel, "Visible", self._gorunur))
            # Kaydetme, bağlantı güncelleme vb. soruları sorma; görünür modda sorun görülebilsin.
            self._tekrarla(lambda: setattr(excel, "DisplayAlerts", self._gorunur))
            self._kitap = self._dosyayi_ac()
            if self._tekrarla(lambda: self._kitap.ReadOnly):
                raise ExcelHatasi(
                    "Dosya salt okunur açıldı. Bu bilgisayarda Excel'de oturum açmış hesabın "
                    "dosyayı düzenleme yetkisi olmalı (Excel > Dosya > Hesap)."
                )
        except ExcelHatasi:
            self.kapat()
            raise
        except Exception as hata:
            self.kapat()
            raise ExcelHatasi(
                f"Excel dosyayı açamadı: {hata}. Adresi ve Excel'de oturum açmış hesabı kontrol edin."
            ) from hata
        return self._kitap

    def _tekrarla(self, islem: Callable, sure: float | None = None):
        """Excel "meşgulüm" derse bekleyip tekrar dener; süre dolarsa vazgeçer."""
        sure = self._mesgul_suresi if sure is None else min(sure, self._mesgul_suresi)
        baslangic = time.monotonic()
        while True:
            try:
                return islem()
            except Exception as hata:
                kod = hata.args[0] if hata.args else None
                if kod not in _EXCEL_MESGUL:
                    raise
                if time.monotonic() - baslangic >= sure:
                    raise ExcelHatasi(
                        f"Excel {sure:.0f} saniye boyunca meşgul kaldı; bir pencere "
                        "(oturum açma, lisans, güncelleme) cevap bekliyor olabilir. ayarlar.toml'a "
                        "excel_gorunur = true yazıp tekrar çalıştırarak Excel'de ne çıktığına bakın."
                    ) from hata
                self._bekle(1)

    def _dosyayi_ac(self):
        """Excel sürümüne göre adresin çözülmüş (boşluklu, Türkçe harfli) ya da kodlu
        (%20, %C3%BC) hâli çalışır; önce çözülmüş hâl denenir."""
        adaylar = self._adres_adaylari()
        for sira, adres in enumerate(adaylar):
            try:
                return self._tekrarla(
                    lambda: self._excel.Workbooks.Open(adres, UpdateLinks=0, ReadOnly=False)
                )
            except ExcelHatasi:
                raise
            except Exception:
                if sira == len(adaylar) - 1:
                    raise

    def _adres_adaylari(self) -> list[str]:
        return list(dict.fromkeys([unquote(self._adres), self._adres]))

    def _sayfa(self, ad: str):
        try:
            kitap = self._ac()
            return self._tekrarla(lambda: kitap.Worksheets(ad))
        except ExcelHatasi:
            raise
        except Exception as hata:
            raise ExcelHatasi(f"'{ad}' sayfası bulunamadı.") from hata

    def _kaydet(self) -> None:
        try:
            self._tekrarla(self._kitap.Save)
        except Exception as hata:
            raise ExcelHatasi(f"Excel dosyası kaydedilemedi: {hata}") from hata
