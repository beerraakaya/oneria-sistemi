"""Komutları uçtan uca, uydurma Excel ve sahte Ollama ile çalıştırır."""

import openpyxl
import pytest

from oneri.__main__ import main

from yardimci import YENI_TARZ_OLUMLU, YENI_TARZ_OLUMSUZ


@pytest.fixture
def ayar_dosyasi(tmp_path, ayarlar):
    yol = tmp_path / "ayarlar.toml"
    yol.write_text(
        "\n".join(
            [
                f"excel_yolu = '{ayarlar.excel_yolu}'",
                f"kurallar_yolu = '{ayarlar.kurallar_yolu}'",
                f"veri_klasoru = '{ayarlar.veri_klasoru}'",
                f"ollama_adresi = '{ayarlar.ollama_adresi}'",
            ]
        ),
        encoding="utf-8",
    )
    return yol


def test_kontrol(ayar_dosyasi, capsys):
    assert main(["--ayarlar", str(ayar_dosyasi), "kontrol"]) == 0
    cikti = capsys.readouterr().out
    assert "Denizli önerisi: 7" in cikti
    assert "Örnek alınabilecek değerlendirme: 5 (yeni tarzda: 2, eski kalıp: 3)" in cikti
    assert "Değerlendirme bekleyen: 1 (satır 9)" in cikti
    assert "qwen2.5: yüklü" in cikti
    assert "Her şey hazır." in cikti


def test_kontrol_eksik_modeli_bildirir(ayar_dosyasi, sahte_ollama, capsys):
    sahte_ollama.modeller = ["qwen2.5:latest"]
    assert main(["--ayarlar", str(ayar_dosyasi), "kontrol"]) == 1
    assert "bge-m3: YÜKLÜ DEĞİL -> ollama pull bge-m3" in capsys.readouterr().out


def test_kor_test_yeni_tarz_ornekleri_cevaplari_gizleyerek_dener(ayar_dosyasi, ayarlar, sahte_ollama, capsys):
    sahte_ollama.sohbet_cevabi = {"gerekce": "Geçerli öneri", "degerlendirme": "Taslak metin."}

    assert main(["--ayarlar", str(ayar_dosyasi), "kor-test"]) == 0

    # Sadece yeni tarzdaki iki öneri test edilir ve hiçbiri kendi cevabını görmez.
    sohbetler = sahte_ollama.sohbetler()
    assert len(sohbetler) == 2
    assert YENI_TARZ_OLUMLU not in sohbetler[0]["messages"][1]["content"]
    assert YENI_TARZ_OLUMSUZ not in sohbetler[1]["messages"][1]["content"]

    cikti = capsys.readouterr().out
    assert "Ekiple aynı karar: 1 / 2 (%50)" in cikti

    raporlar = list(ayarlar.veri_klasoru.glob("kor_test_model_*.xlsx"))
    assert len(raporlar) == 1
    kitap = openpyxl.load_workbook(raporlar[0])
    ozet = {satir[0]: satir[1] for satir in kitap["Özet"].iter_rows(values_only=True)}
    assert ozet["Test edilen öneri"] == 2
    assert ozet["Ekiple aynı karar"] == "1 / 2 (%50)"
    assert ozet["Öneri Değil -> Öneri"] == 1

    tablo = list(kitap["Karşılaştırma"].iter_rows(values_only=True))
    assert tablo[0][:3] == ("Satır", "Tarih", "Konu")
    assert [(s[0], s[5], s[6], s[7]) for s in tablo[1:]] == [
        (6, "Öneri", "Öneri", "Evet"),
        (7, "Öneri Değil", "Öneri", "Hayır"),
    ]
    assert tablo[0][14] == "Yapay Zekânın Gerekçesi"
    assert tablo[1][14] == "Geçerli öneri"


def test_kor_test_adet_ile_sadece_en_yenileri_dener(ayar_dosyasi, sahte_ollama):
    assert main(["--ayarlar", str(ayar_dosyasi), "kor-test", "--adet", "1"]) == 0
    sohbetler = sahte_ollama.sohbetler()
    assert len(sohbetler) == 1
    assert "Kartonların ayrı toplanması." in sohbetler[0]["messages"][1]["content"]


def test_kor_test_gecersiz_cevabi_rapora_yazip_devam_eder(ayar_dosyasi, ayarlar, sahte_ollama):
    sahte_ollama.sohbet_cevabi = "bozuk"
    assert main(["--ayarlar", str(ayar_dosyasi), "kor-test"]) == 0
    rapor = next(ayarlar.veri_klasoru.glob("kor_test_*.xlsx"))
    tablo = list(openpyxl.load_workbook(rapor)["Karşılaştırma"].iter_rows(values_only=True))
    assert len(tablo) == 3
    assert all("geçerli JSON" in satir[13] for satir in tablo[1:])


def test_degerlendir_bekleyen_satir(ayar_dosyasi, capsys):
    assert main(["--ayarlar", str(ayar_dosyasi), "degerlendir", "9"]) == 0
    cikti = capsys.readouterr().out
    assert "Gerekçe      : Geçerli öneri" in cikti
    assert "Onay Durumu  : Öneri" in cikti
    assert "Durum        : Devam Ediyor" in cikti
    assert "Ekibin yazdığı" not in cikti


def test_degerlendir_baska_fabrikanin_satirini_reddeder(ayar_dosyasi, capsys):
    assert main(["--ayarlar", str(ayar_dosyasi), "degerlendir", "5"]) == 1
    assert "5. satırda Denizli önerisi yok" in capsys.readouterr().err


def test_ollama_kapaliyken_anlasilir_hata(tmp_path, ayarlar, capsys):
    yol = tmp_path / "kapali.toml"
    yol.write_text(
        f"excel_yolu = '{ayarlar.excel_yolu}'\nollama_adresi = 'http://127.0.0.1:9'\n", encoding="utf-8"
    )
    assert main(["--ayarlar", str(yol), "kontrol"]) == 1
    assert "Ollama'ya bağlanılamadı" in capsys.readouterr().err


def test_bilinmeyen_ayar(tmp_path, capsys):
    yol = tmp_path / "hatali.toml"
    yol.write_text("dil_model = 'qwen2.5'\n", encoding="utf-8")
    assert main(["--ayarlar", str(yol), "kontrol"]) == 1
    assert "bilinmeyen ayar: dil_model" in capsys.readouterr().err


def test_kor_test_komsu_yontemi_modele_sormaz(ayar_dosyasi, ayarlar, sahte_ollama):
    assert main(["--ayarlar", str(ayar_dosyasi), "kor-test", "--yontem", "komsu"]) == 0
    assert sahte_ollama.sohbetler() == []
    rapor = next(ayarlar.veri_klasoru.glob("kor_test_komsu_*.xlsx"))
    kitap = openpyxl.load_workbook(rapor)
    ozet = {satir[0]: satir[1] for satir in kitap["Özet"].iter_rows(values_only=True)}
    assert "çoğunluk kararı" in ozet["Karar yöntemi"]
    assert ozet["Test edilen öneri"] == 2
    tablo = list(kitap["Karşılaştırma"].iter_rows(values_only=True))
    assert all(satir[14].startswith("Benzer ") for satir in tablo[1:])


def test_kor_test_karma_yontemi(ayar_dosyasi, ayarlar, sahte_ollama):
    assert main(["--ayarlar", str(ayar_dosyasi), "kor-test", "--yontem", "karma"]) == 0
    assert len(sahte_ollama.sohbetler()) == 2
    rapor = next(ayarlar.veri_klasoru.glob("kor_test_karma_*.xlsx"))
    ozet = {s[0]: s[1] for s in openpyxl.load_workbook(rapor)["Özet"].iter_rows(values_only=True)}
    assert ozet["Karar yöntemi"].startswith("karma:")
