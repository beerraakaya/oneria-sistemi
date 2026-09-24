"""SharePoint (Microsoft Graph) bağlantısı, uydurma Excel'i tutan sahte bir Graph sunucusuyla test edilir."""

import json
import re
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from threading import Thread
from urllib.parse import parse_qs, unquote

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from oneri.canli import calistir
from oneri.kaynak import GraphHatasi, SharePointExcel
from oneri.ollama import Ollama
from oneri.taslaklar import TaslakDeposu

SITE = "https://sirket.sharepoint.com/sites/Oneri"
DOSYA_ADRESI = f"{SITE}/Shared Documents/Klasor/Oneri Tablosu.xlsx"


class SahteGraph:
    def __init__(self, excel_yolu):
        self.excel_yolu = excel_yolu
        self.adres = ""
        self.istekler: list[tuple[str, str]] = []
        self.jeton_istekleri: list[dict] = []
        self.hatalar: list[tuple[int, dict]] = []  # sıradaki isteklere verilecek hata cevapları

    def degistir(self, islem):
        kitap = openpyxl.load_workbook(self.excel_yolu)
        islem(kitap["Genel Tablo"])
        kitap.save(self.excel_yolu)


@pytest.fixture
def graph(ornek_excel):
    durum = SahteGraph(ornek_excel)

    class Isleyici(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def _cevap(self, kod, govde=None, basliklar=None):
            # Graph tarihleri sayı olarak döndürür; burada metin olması yeterli.
            veri = govde if isinstance(govde, bytes) else json.dumps(govde or {}, default=str).encode()
            self.send_response(kod)
            for ad, deger in (basliklar or {}).items():
                self.send_header(ad, deger)
            self.send_header("Content-Length", str(len(veri)))
            self.end_headers()
            self.wfile.write(veri)

        def _govde(self):
            return self.rfile.read(int(self.headers.get("Content-Length") or 0))

        def do_POST(self):
            if self.path.startswith("/yanlis/"):
                return self._cevap(400, {"error": "invalid_client", "error_description": "gizli anahtar yanlış"})
            if self.path.endswith("/oauth2/v2.0/token"):
                durum.jeton_istekleri.append(parse_qs(self._govde().decode()))
                return self._cevap(200, {"access_token": "jeton", "expires_in": 3600})
            self._isle("POST", self._govde())

        def do_GET(self):
            self._isle("GET", b"")

        def do_PATCH(self):
            self._isle("PATCH", self._govde())

        def _isle(self, yontem, govde):
            yol = unquote(self.path)
            durum.istekler.append((yontem, yol))
            if self.headers.get("Authorization") != "Bearer jeton":
                return self._cevap(401, {"error": {"message": "yetkisiz"}})
            if durum.hatalar:
                kod, basliklar = durum.hatalar.pop(0)
                return self._cevap(kod, {"error": {"message": f"hata {kod}"}}, basliklar)

            if yol == "/v1.0/sites/sirket.sharepoint.com:/sites/Oneri":
                return self._cevap(200, {"id": "site1"})
            if yol == "/v1.0/sites/site1/drives":
                return self._cevap(200, {"value": [
                    {"id": "d1", "webUrl": f"{SITE}/Shared%20Documents"},
                    {"id": "d2", "webUrl": f"{SITE}/Diger"},
                ]})
            if yol in ("/v1.0/drives/d1/root:/Klasor/Oneri Tablosu.xlsx", "/v1.0/shares/u!paylasim/driveItem") or (
                yol.startswith("/v1.0/shares/u!") and yol.endswith("/driveItem")
            ):
                return self._cevap(200, {"id": "f1", "parentReference": {"driveId": "d1"}})
            if yol == "/v1.0/drives/d1/items/f1/content":
                return self._cevap(200, durum.excel_yolu.read_bytes())

            aralik = re.fullmatch(
                r"/v1\.0/drives/d1/items/f1/workbook/worksheets/Genel Tablo/range\(address='([A-Z0-9:]+)'\)(.*)",
                yol,
            )
            if not aralik:
                return self._cevap(404, {"error": {"message": f"bilinmeyen yol {yol}"}})
            adres, ek = aralik.groups()
            if yontem == "GET" and not ek:
                sayfa = openpyxl.load_workbook(durum.excel_yolu)["Genel Tablo"]
                degerler = [["" if h.value is None else h.value for h in satir] for satir in sayfa[adres]]
                return self._cevap(200, {"values": degerler}, {})
            if yontem == "PATCH" and not ek:
                deger = json.loads(govde)["values"][0][0]
                durum.degistir(lambda s: s.__setitem__(adres, deger))
                return self._cevap(200, {})
            if yontem == "PATCH" and ek == "/format/fill":
                renk = json.loads(govde)["color"].lstrip("#")
                durum.degistir(lambda s: setattr(s[adres], "fill", PatternFill("solid", fgColor=renk)))
                return self._cevap(200, {})
            if yontem == "POST" and ek == "/format/fill/clear":
                durum.degistir(lambda s: setattr(s[adres], "fill", PatternFill(fill_type=None)))
                return self._cevap(200, {})
            return self._cevap(400, {"error": {"message": "beklenmeyen istek"}})

    sunucu = ThreadingHTTPServer(("127.0.0.1", 0), Isleyici)
    Thread(target=sunucu.serve_forever, daemon=True).start()
    durum.adres = f"http://127.0.0.1:{sunucu.server_port}"
    yield durum
    sunucu.shutdown()
    sunucu.server_close()


def _kaynak(graph, adres=DOSYA_ADRESI, **secenekler):
    return SharePointExcel(
        adres,
        "kiraci1",
        "uygulama1",
        "sir1",
        giris_adresi=graph.adres,
        graph_adresi=f"{graph.adres}/v1.0",
        bekle=secenekler.pop("bekle", lambda _: None),
    )


def test_dosya_yoldan_bulunur_ve_indirilir(graph, tmp_path):
    yol = _kaynak(graph).indir(tmp_path / "kopya.xlsx")
    assert yol.read_bytes() == graph.excel_yolu.read_bytes()
    assert [i[1] for i in graph.istekler[:3]] == [
        "/v1.0/sites/sirket.sharepoint.com:/sites/Oneri",
        "/v1.0/sites/site1/drives",
        "/v1.0/drives/d1/root:/Klasor/Oneri Tablosu.xlsx",
    ]
    [jeton] = graph.jeton_istekleri
    assert jeton["client_id"] == ["uygulama1"] and jeton["grant_type"] == ["client_credentials"]
    assert jeton["scope"] == ["https://graph.microsoft.com/.default"]


def test_paylasim_baglantisi_da_kabul_edilir(graph, tmp_path):
    _kaynak(graph, "https://sirket.sharepoint.com/:x:/s/Oneri/EabcXYZ").indir(tmp_path / "k.xlsx")
    assert graph.istekler[0][1].startswith("/v1.0/shares/u!")


def test_satir_okunur_ve_hucreye_yazilir(graph):
    kaynak = _kaynak(graph)
    satir = kaynak.satir_oku("Genel Tablo", 9, 13)
    assert satir[7] == "İş Güvenliği Risk Azaltma" and satir[10:] == [None, None, None]

    kaynak.yaz("Genel Tablo", {"K9": "Metin", "L9": "Öneri"})
    kaynak.boya("Genel Tablo", ["K9"], "#FFF2CC")
    sayfa = openpyxl.load_workbook(graph.excel_yolu)["Genel Tablo"]
    assert (sayfa["K9"].value, sayfa["L9"].value) == ("Metin", "Öneri")
    assert sayfa["K9"].fill.fgColor.rgb.endswith("FFF2CC")


def test_yogunlukta_bekleyip_tekrar_dener(graph):
    beklemeler = []
    graph.hatalar = [(429, {"Retry-After": "7"}), (503, {})]
    kaynak = _kaynak(graph, bekle=beklemeler.append)
    kaynak.yaz("Genel Tablo", {"K9": "Metin"})
    assert beklemeler == [7.0, 2.0]
    assert openpyxl.load_workbook(graph.excel_yolu)["Genel Tablo"]["K9"].value == "Metin"


@pytest.mark.parametrize(
    ("kod", "mesaj"),
    [(403, "erişim izni yok"), (423, "kilitli"), (404, "bulunamadı")],
)
def test_hatalar_anlasilir_mesajla_bildirilir(graph, kod, mesaj):
    graph.hatalar = [(kod, {})]
    with pytest.raises(GraphHatasi, match=mesaj):
        _kaynak(graph).satir_oku("Genel Tablo", 9, 13)


def test_yanlis_kimlik_bilgisi(graph, tmp_path):
    kaynak = SharePointExcel(
        DOSYA_ADRESI, "kiraci1", "uygulama1", "sir1",
        giris_adresi=f"{graph.adres}/yanlis", graph_adresi=f"{graph.adres}/v1.0",
    )
    with pytest.raises(GraphHatasi, match="erişim izni alınamadı.*gizli anahtar yanlış"):
        kaynak.indir(tmp_path / "k.xlsx")


def test_sharepoint_olmayan_adres_reddedilir(graph, tmp_path):
    with pytest.raises(GraphHatasi, match="anlaşılamadı"):
        _kaynak(graph, "https://ornek.com/dosya.xlsx").indir(tmp_path / "k.xlsx")


def test_canli_akis_sharepoint_uzerinden(graph, ayarlar, sahte_ollama):
    from dataclasses import replace

    kaynak = _kaynak(graph)
    ayarlar = replace(ayarlar, excel_yolu=kaynak.indir(ayarlar.veri_klasoru / "sharepoint_kopya.xlsx"))
    depo = TaslakDeposu(ayarlar.taslak_yolu)
    try:
        ozet = calistir(ayarlar, kaynak, Ollama(sahte_ollama.adres), depo, datetime(2026, 9, 24), yaz=lambda _: None)
    finally:
        depo.kapat()

    assert ozet.yazilan == 1
    sayfa = openpyxl.load_workbook(graph.excel_yolu)["Genel Tablo"]
    assert (sayfa["L9"].value, sayfa["M9"].value) == ("Öneri", "Devam Ediyor")
    yazilanlar = [yol for yontem, yol in graph.istekler if yontem == "PATCH" and not yol.endswith("/fill")]
    assert [re.search(r"'(\w+)'", y).group(1) for y in yazilanlar] == ["K9", "L9", "M9"]
