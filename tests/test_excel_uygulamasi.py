"""Excel uygulaması yöntemi, Excel'in Python'dan görünen yüzünü taklit eden sahte nesnelerle test edilir.

Gerçek Excel yalnızca Windows'ta olduğu için burada uydurma Excel dosyasını tutan bir taklit kullanılır.
"""

import shutil
from datetime import datetime

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from oneri.canli import calistir
from oneri.excel import ExcelHatasi
from oneri.kaynak import ExcelUygulamasi
from oneri.ollama import Ollama
from oneri.taslaklar import TaslakDeposu

ADRES = "https://sirket.sharepoint.com/sites/Oneri/Shared%20Documents/Oneri.xlsx"


class SahteIc:
    def __init__(self, hucre):
        self._hucre = hucre

    @property
    def Color(self):
        return self._hucre.fill.fgColor.rgb

    @Color.setter
    def Color(self, bgr):
        # Excel mavi-yeşil-kırmızı sırasıyla verir; dosyaya kırmızı-yeşil-mavi yazılır.
        kod = f"{bgr:06X}"
        self._hucre.fill = PatternFill("solid", fgColor=kod[4:6] + kod[2:4] + kod[0:2])

    @property
    def ColorIndex(self):
        return None

    @ColorIndex.setter
    def ColorIndex(self, deger):
        assert deger == -4142
        self._hucre.fill = PatternFill(fill_type=None)


class SahteAralik:
    def __init__(self, sayfa, adres):
        self._sayfa = sayfa
        self._adres = adres

    @property
    def Value(self):
        if ":" not in self._adres:
            return self._sayfa[self._adres].value
        return tuple(tuple(h.value for h in satir) for satir in self._sayfa[self._adres])

    @Value.setter
    def Value(self, deger):
        assert ":" not in self._adres
        self._sayfa[self._adres] = deger

    @property
    def Interior(self):
        return SahteIc(self._sayfa[self._adres])


class SahteSayfa:
    def __init__(self, sayfa):
        self._sayfa = sayfa

    def Range(self, adres):
        return SahteAralik(self._sayfa, adres)


class SahteKitap:
    def __init__(self, uygulama, yol, salt_okunur):
        self._uygulama = uygulama
        self._yol = yol
        self._kitap = openpyxl.load_workbook(yol)
        self.ReadOnly = salt_okunur

    def Worksheets(self, ad):
        if ad not in self._kitap.sheetnames:
            raise KeyError(ad)
        return SahteSayfa(self._kitap[ad])

    def Save(self):
        self._uygulama.kayit_sayisi += 1
        self._kitap.save(self._yol)

    def SaveCopyAs(self, hedef):
        self._kitap.save(self._yol)
        shutil.copy(self._yol, hedef)

    def Close(self, SaveChanges):
        if SaveChanges:
            self._kitap.save(self._yol)
        self._uygulama.kapanan_kitap += 1


class SahteKitaplar:
    def __init__(self, uygulama):
        self._uygulama = uygulama

    def Open(self, adres, UpdateLinks, ReadOnly):
        self._uygulama.acilan.append(adres)
        if self._uygulama.acilamaz:
            raise OSError("Microsoft Excel dosyaya erişemiyor")
        return SahteKitap(self._uygulama, self._uygulama.yol, self._uygulama.salt_okunur)


class SahteExcel:
    def __init__(self, yol):
        self.yol = yol
        self.Visible = True
        self.DisplayAlerts = True
        self.Workbooks = SahteKitaplar(self)
        self.acilan: list[str] = []
        self.kayit_sayisi = 0
        self.kapanan_kitap = 0
        self.cikti = False
        self.salt_okunur = False
        self.acilamaz = False

    def Quit(self):
        self.cikti = True


@pytest.fixture
def excel(ornek_excel):
    return SahteExcel(ornek_excel)


def _kaynak(excel, adres=ADRES):
    return ExcelUygulamasi(adres, excel_olustur=lambda: excel)


def test_dosya_gorunmeden_acilir_ve_kopyasi_alinir(excel, tmp_path):
    kaynak = _kaynak(excel, ADRES + "?web=1")
    yol = kaynak.indir(tmp_path / "kopya.xlsx")
    assert excel.acilan == [ADRES]  # "?web=1" atılır
    assert (excel.Visible, excel.DisplayAlerts) == (False, False)
    assert openpyxl.load_workbook(yol)["Genel Tablo"]["H9"].value == "İş Güvenliği Risk Azaltma"
    kaynak.kapat()
    assert excel.cikti and excel.kapanan_kitap == 1


def test_satir_okur_yazar_boyar_ve_her_seferinde_kaydeder(excel):
    kaynak = _kaynak(excel)
    satir = kaynak.satir_oku("Genel Tablo", 9, 13)
    assert satir[7] == "İş Güvenliği Risk Azaltma" and satir[10:] == [None, None, None]

    kaynak.yaz("Genel Tablo", {"K9": "Metin", "L9": "Öneri"})
    kaynak.boya("Genel Tablo", ["K9"], "#FFF2CC")
    assert excel.kayit_sayisi == 2
    sayfa = openpyxl.load_workbook(excel.yol)["Genel Tablo"]
    assert (sayfa["K9"].value, sayfa["L9"].value) == ("Metin", "Öneri")
    assert sayfa["K9"].fill.fgColor.rgb.endswith("FFF2CC")

    kaynak.boya("Genel Tablo", ["K9"], None)
    assert openpyxl.load_workbook(excel.yol)["Genel Tablo"]["K9"].fill.fill_type is None
    kaynak.kapat()


def test_salt_okunur_acilirsa_anlasilir_hata_ve_excel_kapanir(excel, tmp_path):
    excel.salt_okunur = True
    with pytest.raises(ExcelHatasi, match="salt okunur"):
        _kaynak(excel).indir(tmp_path / "k.xlsx")
    assert excel.cikti


def test_acilamazsa_anlasilir_hata_ve_excel_kapanir(excel, tmp_path):
    excel.acilamaz = True
    with pytest.raises(ExcelHatasi, match="Excel dosyayı açamadı"):
        _kaynak(excel).indir(tmp_path / "k.xlsx")
    assert excel.cikti


def test_olmayan_sayfa(excel):
    kaynak = _kaynak(excel)
    with pytest.raises(ExcelHatasi, match="'Yok' sayfası bulunamadı"):
        kaynak.satir_oku("Yok", 9, 13)
    kaynak.kapat()


def test_pywin32_yoksa_anlasilir_hata(tmp_path, monkeypatch):
    import builtins

    gercek_import = builtins.__import__

    def sahte_import(ad, *args, **kwargs):
        if ad.startswith("win32com"):
            raise ImportError(ad)
        return gercek_import(ad, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", sahte_import)
    with pytest.raises(ExcelHatasi, match="pywin32"):
        ExcelUygulamasi(ADRES).indir(tmp_path / "k.xlsx")


def test_canli_akis_excel_uygulamasi_uzerinden(excel, ayarlar, sahte_ollama):
    from dataclasses import replace

    kaynak = _kaynak(excel)
    ayarlar = replace(ayarlar, excel_yolu=kaynak.indir(ayarlar.veri_klasoru / "sharepoint_kopya.xlsx"))
    depo = TaslakDeposu(ayarlar.taslak_yolu)
    try:
        ozet = calistir(ayarlar, kaynak, Ollama(sahte_ollama.adres), depo, datetime(2026, 9, 24), yaz=lambda _: None)
    finally:
        depo.kapat()
        kaynak.kapat()

    assert ozet.yazilan == 1
    sayfa = openpyxl.load_workbook(excel.yol)["Genel Tablo"]
    assert (sayfa["L9"].value, sayfa["M9"].value) == ("Öneri", "Devam Ediyor")
    assert sayfa["K9"].fill.fgColor.rgb.endswith("FFF2CC")
    assert excel.cikti
