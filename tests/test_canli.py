"""Canlı akış: uydurma Excel'in bir kopyasına yazma, 7 gün kuralı ve öğrenme."""

from datetime import datetime, timedelta

import openpyxl
import pytest

from oneri.canli import CalismaSuruyor, calistir, tek_calisma
from oneri.kaynak import YerelExcel
from oneri.ollama import Ollama
from oneri.taslaklar import DUZELTILDI, ONAYLANDI, SILINDI, TaslakDeposu

# Uydurma Excel'de Değerlendirme, Onay Durumu ve Durum K, L, M sütunlarında; satır 9 bekliyor.
TASLAK = {"gerekce": "Geçerli öneri", "degerlendirme": "Kayma riskini azaltabilir. Maliyet belirlenmelidir."}
BASLA = datetime(2026, 9, 24, 9, 0)


@pytest.fixture
def depo(ayarlar):
    depo = TaslakDeposu(ayarlar.taslak_yolu)
    yield depo
    depo.kapat()


def _calistir(ayarlar, sahte_ollama, depo, simdi=BASLA, **secenekler):
    ciktilar = []
    ozet = calistir(
        ayarlar,
        secenekler.pop("kaynak", YerelExcel(ayarlar.excel_yolu)),
        Ollama(sahte_ollama.adres),
        depo,
        simdi,
        yaz=ciktilar.append,
        **secenekler,
    )
    return ozet, "\n".join(ciktilar)


def _hucreler(yol, satir):
    sayfa = openpyxl.load_workbook(yol)["Genel Tablo"]
    return [sayfa[f"{s}{satir}"] for s in "KLM"]


def _ekip_yazar(yol, satir, **degerler):
    kitap = openpyxl.load_workbook(yol)
    sayfa = kitap["Genel Tablo"]
    for sutun, deger in degerler.items():
        sayfa[f"{sutun}{satir}"] = deger
    kitap.save(yol)


def test_bekleyen_satiri_doldurur_ve_boyar(ayarlar, sahte_ollama, depo):
    sahte_ollama.sohbet_cevabi = TASLAK
    ozet, cikti = _calistir(ayarlar, sahte_ollama, depo)

    assert ozet.yazilan == 1
    assert "Satır 9" in cikti
    deg, onay, durum = _hucreler(ayarlar.excel_yolu, 9)
    assert (deg.value, onay.value, durum.value) == (TASLAK["degerlendirme"], "Öneri", "Devam Ediyor")
    assert deg.fill.fgColor.rgb.endswith("FFF2CC")

    [kayit] = depo.hepsi()
    assert (kayit.satir, kayit.onay_durumu, kayit.sonuc) == (9, "Öneri", None)


def test_baska_hicbir_hucreye_dokunmaz(ayarlar, sahte_ollama, depo):
    once = [[h.value for h in s] for s in openpyxl.load_workbook(ayarlar.excel_yolu)["Genel Tablo"]]
    _calistir(ayarlar, sahte_ollama, depo)
    sonra = [[h.value for h in s] for s in openpyxl.load_workbook(ayarlar.excel_yolu)["Genel Tablo"]]

    farklar = [
        (r + 1, c) for r, (a, b) in enumerate(zip(once, sonra)) for c, (x, y) in enumerate(zip(a, b)) if x != y
    ]
    assert farklar == [(9, 10), (9, 11), (9, 12)]  # yalnızca satır 9'un K, L, M hücreleri


def test_ikinci_calismada_ayni_satira_yazmaz(ayarlar, sahte_ollama, depo):
    _calistir(ayarlar, sahte_ollama, depo)
    ozet, cikti = _calistir(ayarlar, sahte_ollama, depo, BASLA + timedelta(hours=1))
    assert ozet.yazilan == 0 and "Doldurulacak yeni öneri yok." in cikti
    assert depo.bekleyenler()[0].sonuc is None


def test_yedi_gun_degismeyen_taslak_onaylanir_ve_rengi_kalkar(ayarlar, sahte_ollama, depo):
    _calistir(ayarlar, sahte_ollama, depo)
    ozet, _ = _calistir(ayarlar, sahte_ollama, depo, BASLA + timedelta(days=6, hours=23))
    assert ozet.onaylanan == 0

    # Durum'un sonradan Tamamlandı yapılması süreç ilerlemesidir, düzeltme sayılmaz.
    _ekip_yazar(ayarlar.excel_yolu, 9, M="Tamamlandı")
    ozet, cikti = _calistir(ayarlar, sahte_ollama, depo, BASLA + timedelta(days=7))
    assert ozet.onaylanan == 1 and "onaylandı sayıldı" in cikti
    assert depo.hepsi()[0].sonuc == ONAYLANDI
    assert _hucreler(ayarlar.excel_yolu, 9)[0].fill.fill_type is None


def test_ekibin_duzelttigi_taslak_kaydedilir(ayarlar, sahte_ollama, depo):
    _calistir(ayarlar, sahte_ollama, depo)
    _ekip_yazar(ayarlar.excel_yolu, 9, K="Rutin bakım işidir. Bakım ekibi yenilemelidir.", L="Öneri Değil")

    ozet, _ = _calistir(ayarlar, sahte_ollama, depo, BASLA + timedelta(days=1))
    assert ozet.duzeltilen == 1
    [kayit] = depo.hepsi()
    assert (kayit.sonuc, kayit.ekip_onay_durumu) == (DUZELTILDI, "Öneri Değil")
    assert kayit.ekip_degerlendirme == "Rutin bakım işidir. Bakım ekibi yenilemelidir."


def test_ekip_hucreleri_bosaltirsa_tekrar_doldurmaz(ayarlar, sahte_ollama, depo):
    _calistir(ayarlar, sahte_ollama, depo)
    _ekip_yazar(ayarlar.excel_yolu, 9, K=None, L=None, M=None)
    ozet, _ = _calistir(ayarlar, sahte_ollama, depo, BASLA + timedelta(days=1))
    assert (ozet.duzeltilen, ozet.yazilan) == (1, 0)


def test_silinen_satir(ayarlar, sahte_ollama, depo):
    _calistir(ayarlar, sahte_ollama, depo)
    kitap = openpyxl.load_workbook(ayarlar.excel_yolu)
    kitap["Genel Tablo"].delete_rows(9)
    kitap.save(ayarlar.excel_yolu)

    ozet, _ = _calistir(ayarlar, sahte_ollama, depo, BASLA + timedelta(days=1))
    assert ozet.silinen == 1 and depo.hepsi()[0].sonuc == SILINDI


def test_kontrol_edilmemis_taslak_ornek_alinmaz(ayarlar, sahte_ollama, depo):
    sahte_ollama.sohbet_cevabi = TASLAK
    _calistir(ayarlar, sahte_ollama, depo)
    # Yeni bir öneri gelir; satır 9'daki yapay zekâ metni ona örnek gösterilmemeli.
    kitap = openpyxl.load_workbook(ayarlar.excel_yolu)
    kitap["Genel Tablo"].append(
        [datetime(2026, 9, 25), "Denizli", 1, "Test Kişi", "Üretim", None, None,
         "İş Güvenliği Risk Azaltma", "Rampa zemini kaygan.", "Rampaya kaymaz kaplama yapılması."]
    )
    kitap.save(ayarlar.excel_yolu)

    ozet, _ = _calistir(ayarlar, sahte_ollama, depo, BASLA + timedelta(days=1))
    assert ozet.yazilan == 1
    son_istem = sahte_ollama.sohbetler()[-1]["messages"][1]["content"]
    assert "Rampaya kaymaz kaplama" in son_istem
    assert TASLAK["degerlendirme"] not in son_istem


def test_deneme_hicbir_sey_yazmaz(ayarlar, sahte_ollama, depo):
    sahte_ollama.sohbet_cevabi = TASLAK
    ozet, cikti = _calistir(ayarlar, sahte_ollama, depo, deneme=True)
    assert TASLAK["degerlendirme"] in cikti
    assert [h.value for h in _hucreler(ayarlar.excel_yolu, 9)] == [None, None, None]
    assert depo.hepsi() == [] and ozet.yazilan == 0


class ArayaGirenKaynak(YerelExcel):
    """Program taslağı hazırlarken ekipten birinin satıra yazdığı durumu taklit eder."""

    def satir_oku(self, sayfa, satir, sutun_sayisi):
        hucreler = super().satir_oku(sayfa, satir, sutun_sayisi)
        hucreler[10] = "Ekip yazdı."
        return hucreler


def test_bu_arada_doldurulan_satira_dokunmaz(ayarlar, sahte_ollama, depo):
    ozet, cikti = _calistir(ayarlar, sahte_ollama, depo, kaynak=ArayaGirenKaynak(ayarlar.excel_yolu))
    assert (ozet.yazilan, ozet.atlanan) == (0, 1)
    assert "dokunulmadı" in cikti
    assert [h.value for h in _hucreler(ayarlar.excel_yolu, 9)] == [None, None, None]
    assert depo.hepsi() == []


def test_uretilemeyen_taslak_sonraki_calismaya_kalir(ayarlar, sahte_ollama, depo):
    sahte_ollama.sohbet_cevabi = "bozuk"
    ozet, cikti = _calistir(ayarlar, sahte_ollama, depo)
    assert ozet.hatali == 1 and "tekrar denenecek" in cikti
    assert depo.hepsi() == []


def test_en_fazla_oneri_siniri(ayarlar, sahte_ollama, depo):
    from dataclasses import replace

    ozet, _ = _calistir(replace(ayarlar, en_fazla_oneri=0), sahte_ollama, depo)
    assert (ozet.yazilan, ozet.kalan) == (0, 1)


def test_renk_ayari_bossa_boyamaz(ayarlar, sahte_ollama, depo):
    from dataclasses import replace

    _calistir(replace(ayarlar, taslak_rengi=""), sahte_ollama, depo)
    assert _hucreler(ayarlar.excel_yolu, 9)[0].fill.fill_type is None


def test_ayni_anda_iki_calisma_olmaz(tmp_path):
    with tek_calisma(tmp_path):
        with pytest.raises(CalismaSuruyor):
            with tek_calisma(tmp_path):
                pass
    with tek_calisma(tmp_path):  # ilki bitince kilit kalkar
        pass


def test_eski_kilit_temizlenir(tmp_path):
    (tmp_path / "calisiyor.kilit").write_text("123")
    with tek_calisma(tmp_path, eskime_saniyesi=0):
        pass
    assert not (tmp_path / "calisiyor.kilit").exists()
